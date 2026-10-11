"""Minimal literature survey: one slide (stat tiles, comparison table, takeaways), inserted after 'Existing approaches'.

Numbers are the cited authors' own published results (research/api_agents_literature.md and earlier decks), not our measurements.
The full nine-slide version is kept in survey_slides_full.py.bak."""
from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

AMBER = "E39B2D"
YES_FILL, PARTLY_FILL, NO_FILL = "1B7F8C", "F8DFB9", "E8EDEE"


def build(prs, h) -> list[int]:
    h.prs = prs
    first = len(prs.slides)
    s = h.new_slide(prs, "Literature survey: existing alternatives")
    k = h.SlideKit(s)

    # stat tiles
    tiles = [("9–19%", "enterprise workflows done by clicking agents"),
             ("61%", "complex tasks done by new users (UI-CUBE)"),
             ("14.8% → 29.2%", "WebArena: browsing vs API agent"),
             ("91.1–100%", "VACP-only agents; UI+DOM 28.9–51.1%")]
    for i, (big, txt) in enumerate(tiles):
        x = 0.7 + i * 3.04
        k.card(f"tile{i}", x, 1.3, 2.9, 1.02, fill=h.LIGHT, line=h.LIGHT)
        k.text(f"tile{i}-n", x + 0.14, 1.33, 2.65, 0.4, big, size=21.0, bold=True, color=AMBER if i < 2 else h.TEAL)
        k.text(f"tile{i}-t", x + 0.14, 1.72, 2.65, 0.58, txt, size=15.0, color=h.NAVY2)

    k.text("cmp-label", 0.7, 2.36, 12.0, 0.3, "HOW EXISTING APPROACHES COMPARE", size=15.0, bold=True, color=h.TEAL)
    header = ["Approach", "Operates the app", "No per-app setup", "Knows measures", "Checks the result",
              "Traces figures", "Keeps context"]
    rows = [["Chat-BI, text-to-SQL", "No", "Partly", "Partly", "No", "Partly", "Partly"],
            ["Screen-clicking agents", "Yes", "Yes", "No", "Partly", "No", "Partly"],
            ["API and tool agents", "Partly", "No", "Partly", "Partly", "No", "Partly"],
            ["App-state protocol (VACP)", "Yes", "No", "Partly", "Yes", "Partly", "Yes"],
            ["Ours (API agent + VACP)", "Yes", "No", "Yes", "Yes", "Yes", "Yes"]]
    widths = [3.0] + [1.5] * 6
    gf = h.add_table(s, 0.7, 2.66, widths, [0.7] + [0.42] * 5, [header] + rows, size=15.0, header_size=15.0)
    tbl = gf.table
    for r in range(len(rows) + 1):
        for c in range(len(header)):
            cell = tbl.cell(r, c)
            if c > 0:
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
            if r == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string("0F4C5C")
                continue
            val = rows[r - 1][c]
            if c == 0:
                if r == len(rows):
                    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(AMBER)
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:
                            run.font.color.rgb = RGBColor.from_string("FFFFFF")
                continue
            fill, fg = {"Yes": (YES_FILL, "FFFFFF"), "Partly": (PARTLY_FILL, h.NAVY2), "No": (NO_FILL, h.MUTED)}[val]
            cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(fill)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = RGBColor.from_string(fg)
                    run.font.bold = val == "Yes"

    # takeaways
    k.text("take-label", 0.7, 5.5, 12.0, 0.3, "WHAT THE SURVEY TELLS US", size=15.0, bold=True, color=h.TEAL)
    cards = [("Call the app, not the screen", "Hybrid 38.9% vs browsing 14.8%"),
             ("Check in code", "Ledger and gates: +3.4–7.6 pts"),
             ("Small model reads, code acts", "Gated 7B 0.562 vs 32B 0.300")]
    for i, (t, b) in enumerate(cards):
        x = 0.7 + i * 4.07
        k.card(f"take{i}", x, 5.83, 3.87, 0.7, fill="FFFFFF", line=h.TEAL)
        k.text(f"take{i}-t", x + 0.14, 5.86, 3.6, 0.3, f"{i + 1}  {t}", size=15.0, bold=True, color=h.NAVY)
        k.text(f"take{i}-b", x + 0.14, 6.17, 3.6, 0.32, b, size=15.0, color=h.NAVY2)
    k.text("sv-note", 0.7, 6.58, 12.0, 0.3,
           "Published results, not our measurements; ratings are our reading of the cited work (91 papers reviewed).",
           size=15.0, color=h.MUTED, check=False)
    return list(range(first, len(prs.slides)))
